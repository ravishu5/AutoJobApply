"""
storage/excel_tracker.py — Multi-Sheet Formatted Excel (.xlsx) Exporter & Sync
Generates high-aesthetic, professional Excel workbooks with styled headers,
status color coding, and automatic column sizing.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Optional
from storage.database import Database


HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Slate 800
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10, color="0F172A")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

THIN_BORDER = Border(
    left=Side(style='thin', color="E2E8F0"),
    right=Side(style='thin', color="E2E8F0"),
    top=Side(style='thin', color="E2E8F0"),
    bottom=Side(style='thin', color="E2E8F0")
)


def _style_header_row(ws, headers: List[str]):
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def _auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(50, max(12, max_len + 3))


def export_tracker_xlsx(output_path: str = "data/Job_Hunt_Tracker.xlsx", db: Optional[Database] = None) -> str:
    """Export complete relational data from SQLite into a formatted multi-tab Excel file."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    if not db:
        db = Database()

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # 1. Sheet: All Discovered Jobs
    # -------------------------------------------------------------
    ws_jobs = wb.create_sheet(title="All Jobs")
    job_headers = ["Job ID", "Job Title", "Company", "Location", "Source", "Match Score (%)", "Status", "Date Posted", "Job URL"]
    _style_header_row(ws_jobs, job_headers)

    jobs = db.get_all_jobs()
    for j in jobs:
        score_val = j.get("match_score", 0.0)
        row = [
            j.get("id", ""),
            j.get("title", ""),
            j.get("company", ""),
            j.get("location", ""),
            j.get("source", ""),
            f"{score_val:.1f}%",
            j.get("status", "new").upper(),
            j.get("date_posted", ""),
            j.get("job_url", "")
        ]
        ws_jobs.append(row)

        curr_row = ws_jobs.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws_jobs.cell(row=curr_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            # Highlight status column
            if col_idx == 7:
                status_val = str(cell.value).upper()
                if "APPLIED" in status_val or "SUBMITTED" in status_val:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # green
                elif "SHORTLISTED" in status_val:
                    cell.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid") # blue
                elif "NEW" in status_val:
                    cell.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # yellow

    _auto_fit_columns(ws_jobs)

    # -------------------------------------------------------------
    # 2. Sheet: Applications
    # -------------------------------------------------------------
    ws_apps = wb.create_sheet(title="Applications")
    app_headers = ["App ID", "Job ID", "Company", "Role", "Platform", "Applied At", "Status", "Notes"]
    _style_header_row(ws_apps, app_headers)

    apps = db.get_all_applications()
    for a in apps:
        row = [
            a.get("id", ""),
            a.get("job_id", ""),
            a.get("company", ""),
            a.get("title", ""),
            a.get("apply_method", ""),
            a.get("applied_at", "")[:19].replace("T", " "),
            a.get("status", "pending").upper(),
            a.get("notes", "")
        ]
        ws_apps.append(row)
        curr_row = ws_apps.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws_apps.cell(row=curr_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    _auto_fit_columns(ws_apps)

    # -------------------------------------------------------------
    # 3. Sheet: LinkedIn Post Leads
    # -------------------------------------------------------------
    ws_leads = wb.create_sheet(title="LinkedIn Post Leads")
    lead_headers = ["Post ID", "Author / Recruiter", "Company", "Target Role", "Direct Email", "Outreach Status", "Date Found", "Post URL"]
    _style_header_row(ws_leads, lead_headers)

    posts = db.get_all_linkedin_posts()
    for p in posts:
        row = [
            p.get("id", ""),
            p.get("author", ""),
            p.get("company", ""),
            p.get("role", ""),
            p.get("email", ""),
            p.get("outreach_status", "").upper(),
            p.get("date_found", ""),
            p.get("post_url", "")
        ]
        ws_leads.append(row)
        curr_row = ws_leads.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws_leads.cell(row=curr_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    _auto_fit_columns(ws_leads)

    # -------------------------------------------------------------
    # 4. Sheet: Referral Network
    # -------------------------------------------------------------
    ws_refs = wb.create_sheet(title="Referral Network")
    ref_headers = ["Name", "Headline", "Company", "LinkedIn Profile", "300-Char Connection Note Draft", "Status", "Date Added"]
    _style_header_row(ws_refs, ref_headers)

    refs = db.get_all_referral_contacts()
    for r in refs:
        row = [
            r.get("name", ""),
            r.get("headline", ""),
            r.get("company", ""),
            r.get("linkedin_url", ""),
            r.get("connection_note", ""),
            r.get("status", "").upper(),
            r.get("date_added", "")
        ]
        ws_refs.append(row)
        curr_row = ws_refs.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws_refs.cell(row=curr_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
    _auto_fit_columns(ws_refs)

    # -------------------------------------------------------------
    # 5. Sheet: Metrics & Dashboard
    # -------------------------------------------------------------
    ws_kpi = wb.create_sheet(title="Dashboard KPIs")
    _style_header_row(ws_kpi, ["Metric Key", "Value", "Description"])

    metrics = db.get_metrics()
    kpi_rows = [
        ("Total Jobs Discovered", metrics.get("total_jobs", 0), "All jobs aggregated across LinkedIn, Indeed, Glassdoor, etc."),
        ("Shortlisted Opportunities", metrics.get("shortlisted_jobs", 0), "Jobs scoring above candidate's match threshold"),
        ("Applications Submitted", metrics.get("applied_jobs", 0), "Jobs where application was successfully automated"),
        ("LinkedIn Post Leads", metrics.get("linkedin_leads", 0), "Hiring announcements found directly from recruiter feeds"),
        ("Referral Targets Found", metrics.get("referral_contacts", 0), "Senior engineers and recruiters mapped for referral reach"),
        ("Average Match Score", f"{metrics.get('avg_match_score', 0.0)}%", "Average skill/ATS alignment across discovered jobs")
    ]
    for kpi, val, desc in kpi_rows:
        ws_kpi.append([kpi, val, desc])
        curr_row = ws_kpi.max_row
        ws_kpi.cell(row=curr_row, column=1).font = BOLD_FONT
        ws_kpi.cell(row=curr_row, column=2).font = BOLD_FONT
        ws_kpi.cell(row=curr_row, column=3).font = REGULAR_FONT
        for col_idx in range(1, 4):
            ws_kpi.cell(row=curr_row, column=col_idx).border = THIN_BORDER
    _auto_fit_columns(ws_kpi)

    wb.save(output_path)
    return output_path
