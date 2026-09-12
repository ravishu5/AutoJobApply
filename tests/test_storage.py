"""
tests/test_storage.py — Unit Tests for Database and Excel Tracker
"""

import unittest
import os
import openpyxl
from storage.database import Database
from storage.excel_tracker import export_tracker_xlsx


class TestStorageSubsystem(unittest.TestCase):
    def setUp(self):
        self.test_db_path = "data/test_job_hunt.db"
        self.test_excel_path = "data/test_tracker.xlsx"
        if os.path.exists(self.test_db_path):
            os.remove(self.test_db_path)
        if os.path.exists(self.test_excel_path):
            os.remove(self.test_excel_path)
        self.db = Database(self.test_db_path)

    def tearDown(self):
        if os.path.exists(self.test_db_path):
            os.remove(self.test_db_path)
        if os.path.exists(self.test_excel_path):
            os.remove(self.test_excel_path)

    def test_database_crud_and_metrics(self):
        job = {
            "id": "job_12345",
            "title": "Senior Python Architect",
            "company": "Apex Cloud Systems",
            "location": "Remote",
            "job_url": "https://example.com/jobs/12345",
            "description": "Building scalable Python platforms.",
            "match_score": 88.5,
            "matched_keywords": ["python", "aws", "docker"],
            "missing_keywords": ["k8s"],
            "status": "shortlisted",
            "source": "linkedin"
        }
        self.db.upsert_job(job)

        jobs = self.db.get_all_jobs(min_score=80.0)
        self.assertEqual(len(jobs), 1)
        self.assertEqual(jobs[0]["company"], "Apex Cloud Systems")
        self.assertEqual(jobs[0]["matched_keywords"], ["python", "aws", "docker"])

        # Record application
        self.db.record_application(
            job_id="job_12345",
            company="Apex Cloud Systems",
            title="Senior Python Architect",
            applied_to="https://example.com/jobs/12345",
            apply_method="linkedin_easy_apply",
            status="submitted",
            notes="Submitted via automated stepper"
        )

        apps = self.db.get_all_applications()
        self.assertEqual(len(apps), 1)
        self.assertEqual(apps[0]["status"], "submitted")

        # Check updated metrics
        metrics = self.db.get_metrics()
        self.assertEqual(metrics["total_jobs"], 1)
        self.assertEqual(metrics["applied_jobs"], 1)

    def test_excel_export_multi_sheet(self):
        # Seed test data
        self.db.upsert_job({
            "id": "job_test_1",
            "title": "AI Platform Engineer",
            "company": "OpenAI Partner",
            "location": "Remote",
            "job_url": "https://example.com/ai",
            "match_score": 92.0,
            "status": "shortlisted",
            "source": "indeed"
        })

        self.db.insert_linkedin_post({
            "author": "Recruiter Jane",
            "company": "ScaleTech",
            "role": "Python Dev",
            "snippet": "We are hiring 3 python developers. Email jane@scaletech.io",
            "email": "jane@scaletech.io",
            "post_url": "https://linkedin.com/posts/jane-1"
        })

        self.db.insert_referral_contact({
            "name": "David Smith",
            "headline": "Engineering Manager at Stripe",
            "company": "Stripe",
            "linkedin_url": "https://linkedin.com/in/davidsmith",
            "connection_note": "Hi David, would love to connect!"
        })

        out_path = export_tracker_xlsx(self.test_excel_path, self.db)
        self.assertTrue(os.path.exists(out_path))

        wb = openpyxl.load_workbook(out_path)
        expected_sheets = ["All Jobs", "Applications", "LinkedIn Post Leads", "Referral Network", "Dashboard KPIs"]
        for sheet_name in expected_sheets:
            self.assertIn(sheet_name, wb.sheetnames)

        ws_jobs = wb["All Jobs"]
        self.assertGreaterEqual(ws_jobs.max_row, 2)
        self.assertEqual(ws_jobs.cell(row=2, column=2).value, "AI Platform Engineer")


if __name__ == "__main__":
    unittest.main()
